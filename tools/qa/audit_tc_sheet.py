# -*- coding: utf-8 -*-
"""WeaponHero_TC.xlsx 의 사전 조건·테스트 스텝 열을 점검한다.

apply_test_steps.py 를 돌린 뒤 이것으로 확인한다. 보는 것은 여덟이다.
초안과 글자까지 같은가 · 번호가 1부터 이어지는가 · 작업용 표시나 마크다운이
새지 않았는가 · 사전 조건에 행동이 섞이지 않았는가 · 스텝이 복붙으로 겹치지
않았는가 · 문자가 유실되지 않았는가 · 기대 결과와 결과가 비지 않았는가 ·
읽었을 때 어색한 자리가 없는가.

    uv run --python 3.12 --with openpyxl python tools/qa/audit_tc_sheet.py

알려진 오탐 둘. TC-EDGE-27 의 1600x900 은 "숫자 붙음"으로 잡히고,
N/A 둘(TC-EDGE-22 · TC-EDGE-17)은 빈 칸이 "빈 줄"로 잡힌다. 전부 정상이다.
"""
import io, sys, os, re, importlib.util
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

spec = importlib.util.spec_from_file_location("m", r"tools\qa\apply_test_steps.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
draft = m.parse_draft(m.DRAFT.read_text(encoding="utf-8"))
expects = m.parse_expects(m.CASES.read_text(encoding="utf-8"))

wb = openpyxl.load_workbook(r"docs\qa\WeaponHero_TC.xlsx")
ws = wb["TC 목록"]
ID = re.compile(r"TC-[A-Z]+-\d+")
issues = []
def bad(tc, why, detail=""):
    issues.append((tc, why, detail))

seen_steps = {}
for r in range(2, ws.max_row + 1):
    d = str(ws.cell(r, 4).value or "")
    mm = ID.search(d)
    if not mm:
        bad("행%d" % r, "TC ID 없음", d[:40]); continue
    tc = mm.group(0)
    pre = ws.cell(r, 5).value or ""
    st  = ws.cell(r, 6).value or ""
    exp = ws.cell(r, 7).value or ""
    res = ws.cell(r, 8).value or ""

    # N/A 는 절차가 없는 것이 정상이다. TC-EDGE-17 은 2회전에서 내려왔다 —
    # 선행 조건(적 생존 · 출구 열림)이 동시에 성립하지 않아 절차를 쓸 수 없다.
    if tc in ("TC-EDGE-22", "TC-EDGE-17"):
        if pre or st: bad(tc, "N/A 인데 칸이 차 있다", repr(pre)+repr(st))
        continue

    # 0) 기대 결과가 TC 문서와 같은가 — 엑셀은 파생물이다
    want = expects.get(tc)
    if want and exp.strip() != want:
        bad(tc, "기대 결과가 TC 문서와 다르다", "%r vs %r" % (exp, want))

    # 1) 초안과 글자까지 같은가
    dpre, dsteps = draft.get(tc, ("", []))
    if st != "\n".join(dsteps): bad(tc, "스텝이 초안과 다르다")
    if pre != dpre: bad(tc, "사전조건이 초안과 다르다", "%r vs %r" % (pre, dpre))

    # 2) 스텝 형식
    lines = st.split("\n")
    nums = [l for l in lines if re.match(r"^\d+\.", l)]
    if not lines[0].startswith("1."): bad(tc, "1. 로 시작하지 않는다", lines[0][:40])
    got = [int(re.match(r"^(\d+)\.", l).group(1)) for l in nums]
    if got != list(range(1, len(got)+1)): bad(tc, "번호가 이어지지 않는다", str(got))
    if len(nums) < 2: bad(tc, "스텝이 1줄뿐", st[:40])

    # 3) 마크다운·표시 잔재
    for tok in ("**", "`", "[확인]", "[에디터]", "[자동]", "~~"):
        if tok in st or tok in pre: bad(tc, "잔재 %s" % tok, (st+pre)[:40])

    # 4) 사전조건에 행동이 섞였나
    for v in ("누른다", "본다", "실행한다", "이동한다", "때린다", "고른다", "적는다"):
        if v in pre: bad(tc, "사전조건에 행동", pre)

    # 5) 스텝 중복 (복붙 사고)
    seen_steps.setdefault(st, []).append(tc)

    # 6) 문자 유실
    if "099" in pre or "099" in st: bad(tc, "물결표 유실 의심", pre or st[:40])
    if re.search(r"\d{4,}", pre): bad(tc, "숫자 붙음 의심", pre)

    # 7) 결과·기대결과가 비었나
    if not exp: bad(tc, "기대 결과 빈칸")
    if not res: bad(tc, "결과 빈칸")

for st, tcs in seen_steps.items():
    if len(tcs) > 1: bad(", ".join(tcs), "스텝이 완전히 같다", st[:50])

print("=== 감사 결과: 지적 %d건 ===" % len(issues))
for tc, why, detail in issues:
    print("  %-26s %-24s %s" % (tc, why, detail))

# --- 2차 점검: 사람이 읽었을 때 어색한 자리 ---
print("\n=== 문장 점검 ===")
warn = 0
for r in range(2, ws.max_row + 1):
    d = str(ws.cell(r,4).value or ""); mm = ID.search(d)
    if not mm: continue
    tc = mm.group(0)
    pre = str(ws.cell(r,5).value or ""); st = str(ws.cell(r,6).value or "")
    for w, why in [
        ("에디터 수행 H", "구분점 빠짐"), ("· 또는", "접속이 깨짐"),
        ("  ", "이중 공백"), (" ·$", "꼬리 구분점"),
    ]:
        if re.search(w, pre) or re.search(w, st):
            print("  %-12s %-12s %s" % (tc, why, pre or st[:40])); warn += 1
    if pre.startswith("·") or pre.endswith("·"): print("  %-12s 구분점 끝" % tc); warn += 1
    if st.count("\n") + 1 != len([l for l in st.split("\n") if l.strip()]):
        print("  %-12s 빈 줄 포함" % tc); warn += 1
print("  없음" if not warn else "  지적 %d" % warn)
