# -*- coding: utf-8 -*-
"""어두운 덱을 어두운 채로 한 단계 밝힌다.

라이트 테마로 뒤집지 않는다. 배경(#14141F 계열)과 카드를 명도 +0.04,
글자와 선을 +0.05 올려 **전체를 한 칸 들어올린다.** 색상(H)과 채도(S)는
그대로라 주황·초록 강조와 스크린샷이 지금 인상을 유지한다.

배경을 밝히면 밝은 글자의 대비는 오히려 떨어진다. 그래서 글자를 배경보다
더 많이 올려 모든 조합이 지금보다 같거나 높게 끝나도록 값을 맞췄다.
초록은 명도를 올리면 채도가 죽어 +0.07, 크림은 이미 흰색에 가까워 +0.05 다.

주의 — 같은 색이 역할에 따라 다르게 쓰인다.
  #14141F  배경 채움 17개 · 주황 카드 위 글자 23런
           채움만 밝히고 글자는 그대로 둔다. 주황이 밝아지므로 이 글자의
           대비는 7.92 에서 8.88 로 저절로 올라간다.

    uv run --python 3.12 --with python-pptx python tools/deck/lighten_theme.py
    ... --dry   바꾸지 않고 무엇이 바뀔지만 본다
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_FILL

DECK = Path.home() / "Desktop" / "손주환 2차 개인프로젝트 QA 포트폴리오.pptx"

# 채움 — 배경과 카드. 어두운 순서(14141F < 1A1A28 < 1E1E2C < 20202F)를 지킨다
FILL = {
    "14141F": "1C1C2B",   # 슬라이드 배경판
    "1A1A28": "222234",
    "1E1E2C": "262638",
    "20202F": "28283B",   # 가장 밝은 카드 · 표 칸
    "FF8C4D": "FF9C67",   # 주황 카드 · 구분 막대
    "FF5A6E": "FF6F81",
}

# 선 — 배경이 밝아진 만큼 같이 올려야 구분선이 안 묻힌다
LINE = {
    "3A3852": "454261",
    "5D5B6F": "69667D",
    "918EA6": "9F9CB1",
    "FF8C4D": "FF9C67",
    "FF5A6E": "FF6F81",
}

# 글자 — 배경 상승분(+0.04)보다 크게(+0.05) 올린다. #14141F 은 주황 위
# 어두운 글자라 그대로 둔다
TEXT = {
    "EDEAE4": "F7F6F3",
    "D9D9D9": "E6E6E6",
    "D8D8D8": "E5E5E5",
    "C7C7C9": "D4D4D6",
    "BEBEC8": "CCCCD4",
    "BCBBC9": "CAC9D4",
    "A8A5BC": "B6B4C7",
    "A6A4B7": "B4B2C2",
    "918EA6": "9F9CB1",
    "6ECB63": "87D47E",   # 초록만 +0.07 — 명도를 올리면 채도가 죽는다
    "FF8C4D": "FF9C67",
    "FF8FA0": "FFA9B6",
}


def main(argv):
    dry = "--dry" in argv
    prs = Presentation(DECK)
    n = {"채움": 0, "표 칸": 0, "선": 0, "글자": 0, "그대로 둔 글자": 0}

    def set_fill(obj, table=False):
        try:
            if obj.fill.type != MSO_FILL.SOLID:
                return
            cur = str(obj.fill.fore_color.rgb).upper()
        except Exception:
            return
        if cur in FILL:
            n["표 칸" if table else "채움"] += 1
            if not dry:
                obj.fill.fore_color.rgb = RGBColor.from_string(FILL[cur])

    for slide in prs.slides:
        for shape in slide.shapes:
            set_fill(shape)
            try:
                if shape.line.fill.type == MSO_FILL.SOLID:
                    cur = str(shape.line.color.rgb).upper()
                    if cur in LINE:
                        n["선"] += 1
                        if not dry:
                            shape.line.color.rgb = RGBColor.from_string(LINE[cur])
            except Exception:
                pass

            cells = []
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        set_fill(cell, table=True)
                        cells.append(cell.text_frame)
            frames = cells + ([shape.text_frame] if shape.has_text_frame else [])

            for tf in frames:
                for para in tf.paragraphs:
                    for run in para.runs:
                        if not run.text.strip():
                            continue
                        color = run.font.color
                        try:
                            if color is None or color.type is None or color.rgb is None:
                                continue
                            cur = str(color.rgb).upper()
                        except Exception:
                            continue
                        if cur == "14141F":       # 주황 카드 위 어두운 글자
                            n["그대로 둔 글자"] += 1
                            continue
                        if cur in TEXT:
                            n["글자"] += 1
                            if not dry:
                                color.rgb = RGBColor.from_string(TEXT[cur])

    for k, v in n.items():
        print("  %-12s %4d" % (k, v))
    if dry:
        print("\n--dry 라서 저장하지 않았다.")
    else:
        prs.save(DECK)
        print("\n저장했다: %s" % DECK.name)


if __name__ == "__main__":
    main(sys.argv[1:])
